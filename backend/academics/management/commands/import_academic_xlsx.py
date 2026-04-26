from __future__ import annotations

from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from academics.models import (
    Carrera,
    Competencia,
    EspacioCurricular,
    PlanEstudio,
    PlanEstudioEC,
    UnidadAcademica,
)
from planning.models import AsignacionDocente


User = get_user_model()


class DryRunRollback(Exception):
    """Internal exception to rollback dry-run imports."""


class Command(BaseCommand):
    help = (
        "Importa datos académicos iniciales desde un archivo XLSX con hojas "
        "Propuestas, Espacios, Competencias y Docentes."
    )

    def add_arguments(self, parser):
        parser.add_argument("xlsx_path", help="Ruta al archivo XLSX")
        parser.add_argument(
            "--unidad-sigla",
            required=True,
            help="Sigla de la unidad académica a crear o reutilizar",
        )
        parser.add_argument(
            "--unidad-nombre",
            required=True,
            help="Nombre de la unidad académica a crear o reutilizar",
        )
        parser.add_argument(
            "--career-level",
            default=Carrera.DegreeLevel.GRADO,
            choices=[choice for choice, _ in Carrera.DegreeLevel.choices],
            help="Nivel por defecto para las carreras importadas",
        )
        parser.add_argument(
            "--plan-credits",
            required=True,
            type=int,
            help="Cantidad de créditos a asignar a cada plan importado",
        )
        parser.add_argument(
            "--vigente-desde",
            required=True,
            help="Fecha de vigencia inicial para los planes (YYYY-MM-DD)",
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Procesa el archivo sin persistir cambios",
        )

    def handle(self, *args, **options):
        workbook_path = Path(options["xlsx_path"]).expanduser().resolve()
        if not workbook_path.exists():
            raise CommandError(f"No existe el archivo: {workbook_path}")

        try:
            vigente_desde = timezone.datetime.fromisoformat(options["vigente_desde"]).date()
        except ValueError as exc:
            raise CommandError("--vigente-desde debe tener formato YYYY-MM-DD") from exc

        self.stats = defaultdict(lambda: {"creados": 0, "actualizados": 0, "sin_cambios": 0})
        self.warnings: list[str] = []

        workbook = load_workbook(filename=workbook_path, data_only=True)
        required_sheet = "Propuestas"
        if required_sheet not in workbook.sheetnames:
            raise CommandError("El archivo XLSX debe contener la hoja 'Propuestas'.")

        plan_by_propuesta_id: dict[str, PlanEstudio] = {}
        plan_by_code: dict[str, PlanEstudio] = {}

        try:
            with transaction.atomic():
                unidad_academica = self._upsert_unidad_academica(
                    sigla=options["unidad_sigla"].strip(),
                    nombre=options["unidad_nombre"].strip(),
                )
                plan_by_propuesta_id, plan_by_code = self._import_propuestas(
                    workbook[required_sheet],
                    unidad_academica=unidad_academica,
                    career_level=options["career_level"],
                    plan_credits=options["plan_credits"],
                    vigente_desde=vigente_desde,
                )
                if "Espacios" in workbook.sheetnames:
                    self._import_espacios(workbook["Espacios"], plan_by_propuesta_id)
                else:
                    self.warnings.append("No se encontró la hoja 'Espacios'; se omitió la importación de espacios curriculares.")

                if "Competencias" in workbook.sheetnames:
                    self._import_competencias(workbook["Competencias"], plan_by_code)
                else:
                    self.warnings.append("No se encontró la hoja 'Competencias'; se omitió la importación de competencias.")

                if "Docentes" in workbook.sheetnames:
                    self._import_docentes(workbook["Docentes"], vigente_desde=vigente_desde)
                else:
                    self.warnings.append("No se encontró la hoja 'Docentes'; se omitió la importación de docentes.")

                if options["dry_run"]:
                    raise DryRunRollback()
        except DryRunRollback:
            self.stdout.write(self.style.WARNING("Dry-run ejecutado: los cambios fueron revertidos."))

        self._print_summary()

    def _sheet_rows(self, worksheet):
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [self._clean_string(value) for value in rows[0]]
        normalized_headers = [header if header else f"__empty_{index}" for index, header in enumerate(headers)]
        mapped_rows = []
        for row in rows[1:]:
            if not any(cell not in (None, "") for cell in row):
                continue
            mapped_rows.append({
                normalized_headers[index]: row[index] if index < len(row) else None
                for index in range(len(normalized_headers))
            })
        return mapped_rows

    def _upsert_unidad_academica(self, *, sigla: str, nombre: str) -> UnidadAcademica:
        unidad, created = UnidadAcademica.objects.get_or_create(
            sigla=sigla,
            defaults={"nombre": nombre},
        )
        if created:
            self.stats["unidad_academica"]["creados"] += 1
            return unidad

        if unidad.nombre != nombre:
            unidad.nombre = nombre
            unidad.save(update_fields=["nombre"])
            self.stats["unidad_academica"]["actualizados"] += 1
        else:
            self.stats["unidad_academica"]["sin_cambios"] += 1
        return unidad

    def _import_propuestas(
        self,
        worksheet,
        *,
        unidad_academica: UnidadAcademica,
        career_level: str,
        plan_credits: int,
        vigente_desde,
    ):
        required_columns = {"id_propuesta", "estado", "nombre", "nombre_plan", "titulo", "normativa"}
        rows = self._sheet_rows(worksheet)
        if not rows:
            raise CommandError("La hoja 'Propuestas' está vacía.")

        missing_columns = required_columns - set(rows[0].keys())
        if missing_columns:
            missing = ", ".join(sorted(missing_columns))
            raise CommandError(f"Faltan columnas requeridas en 'Propuestas': {missing}")

        plan_by_propuesta_id: dict[str, PlanEstudio] = {}
        plan_by_code: dict[str, PlanEstudio] = {}

        for row_number, row in enumerate(rows, start=2):
            career_name = self._require_value(row, "nombre", row_number, "Propuestas")
            career_code = self._normalize_code(self._require_value(row, "nombre_plan", row_number, "Propuestas"))
            plan_display = self._require_value(row, "titulo", row_number, "Propuestas")
            normativa = self._require_value(row, "normativa", row_number, "Propuestas")
            proposal_id = self._normalize_code(self._require_value(row, "id_propuesta", row_number, "Propuestas"))

            carrera = self._upsert_carrera(
                codigo=career_code,
                nombre=career_name,
                unidad_academica=unidad_academica,
                nivel=career_level,
            )
            plan = self._upsert_plan_estudio(
                carrera=carrera,
                nombre=career_code,
                descripcion=plan_display,
                ordenanza=normativa,
                creditos=plan_credits,
                vigente_desde=vigente_desde,
                vigente_hasta=self._infer_end_date(row.get("estado")),
            )
            plan_by_propuesta_id[proposal_id] = plan
            plan_by_code[career_code] = plan
            for alias in self._extract_plan_aliases(plan_display):
                plan_by_code[alias] = plan

        return plan_by_propuesta_id, plan_by_code

    def _import_espacios(self, worksheet, plan_by_propuesta_id: dict[str, PlanEstudio]):
        required_columns = {
            "id_propuesta",
            "codigo_materia",
            "nombre_materia",
            "año_materia",
            "cuatrimestre_materia",
            "horas_i_p",
            "horas_f_p",
            "horas_t_a_e",
            "creditos",
        }
        rows = self._sheet_rows(worksheet)
        if not rows:
            self.warnings.append("La hoja 'Espacios' está vacía; no se importaron espacios curriculares.")
            return

        missing_columns = required_columns - set(rows[0].keys())
        if missing_columns:
            missing = ", ".join(sorted(missing_columns))
            raise CommandError(f"Faltan columnas requeridas en 'Espacios': {missing}")

        for row_number, row in enumerate(rows, start=2):
            proposal_id = self._normalize_code(self._require_value(row, "id_propuesta", row_number, "Espacios"))
            plan = plan_by_propuesta_id.get(proposal_id)
            if plan is None:
                raise CommandError(
                    f"Hoja 'Espacios', fila {row_number}: id_propuesta '{proposal_id}' no existe en 'Propuestas'."
                )

            codigo = self._normalize_code(self._require_value(row, "codigo_materia", row_number, "Espacios"))
            nombre = self._require_value(row, "nombre_materia", row_number, "Espacios")
            horas_ip = self._to_int(row.get("horas_i_p")) + self._to_int(row.get("horas_f_p"))
            horas_ta = self._to_int(row.get("horas_t_a_e"))
            creditos = self._to_int(row.get("creditos"))
            anio_cursada = self._extract_year(row.get("año_materia"), row_number)
            periodo = self._infer_periodo(row.get("cuatrimestre_materia"), row.get("bimestre_materia"))
            tipo_espacio = self._infer_tipo_espacio(horas_ip, horas_ta)

            espacio_curricular = self._upsert_espacio_curricular(
                codigo=codigo,
                nombre=nombre,
                anio_cursada=anio_cursada,
                periodo=periodo,
                creditos=creditos,
                horas_ip=horas_ip,
                horas_ta=horas_ta,
                tipo_espacio=tipo_espacio,
            )
            self._upsert_plan_estudio_ec(plan=plan, espacio=espacio_curricular)

    def _import_competencias(self, worksheet, plan_by_code: dict[str, PlanEstudio]):
        required_columns = {"titulo", "codigo", "descripcion", "tipo"}
        rows = self._sheet_rows(worksheet)
        if not rows:
            self.warnings.append("La hoja 'Competencias' está vacía; no se importaron competencias.")
            return

        missing_columns = required_columns - set(rows[0].keys())
        if missing_columns:
            missing = ", ".join(sorted(missing_columns))
            raise CommandError(f"Faltan columnas requeridas en 'Competencias': {missing}")

        for row_number, row in enumerate(rows, start=2):
            plan_code = self._normalize_code(self._require_value(row, "titulo", row_number, "Competencias"))
            plan = plan_by_code.get(plan_code)
            if plan is None:
                raise CommandError(
                    f"Hoja 'Competencias', fila {row_number}: titulo '{plan_code}' no coincide con un plan importado."
                )

            codigo = self._require_value(row, "codigo", row_number, "Competencias")
            descripcion = self._require_value(row, "descripcion", row_number, "Competencias")
            tipo = self._clean_string(row.get("tipo"))
            nombre = descripcion[:255]
            if tipo:
                nombre = f"{tipo}: {nombre}"[:255]

            self._upsert_competencia(
                plan=plan,
                codigo=codigo,
                nombre=nombre,
                descripcion=descripcion,
            )

    def _import_docentes(self, worksheet, *, vigente_desde):
        required_columns = {"codigo", "apellido", "nombres", "nombre-3"}
        rows = self._sheet_rows(worksheet)
        if not rows:
            self.warnings.append("La hoja 'Docentes' está vacía; no se importaron docentes.")
            return

        missing_columns = required_columns - set(rows[0].keys())
        if missing_columns:
            missing = ", ".join(sorted(missing_columns))
            raise CommandError(f"Faltan columnas requeridas en 'Docentes': {missing}")

        for row_number, row in enumerate(rows, start=2):
            espacio_codigo = self._normalize_code(self._require_value(row, "codigo", row_number, "Docentes"))
            espacio = EspacioCurricular.objects.filter(codigo=espacio_codigo).first()
            if espacio is None:
                self.warnings.append(
                    f"Hoja 'Docentes', fila {row_number}: codigo '{espacio_codigo}' no existe en EspacioCurricular; fila omitida."
                )
                continue

            categoria = self._map_docente_categoria(row.get("nombre-3"), row_number)
            docente = self._upsert_docente_user(row, row_number)
            self._upsert_asignacion_docente(
                docente=docente,
                espacio=espacio,
                categoria=categoria,
                vigente_desde=vigente_desde,
            )

    def _upsert_carrera(self, *, codigo: str, nombre: str, unidad_academica: UnidadAcademica, nivel: str) -> Carrera:
        carrera, created = Carrera.objects.get_or_create(
            codigo=codigo,
            defaults={
                "nombre": nombre,
                "unidad_academica": unidad_academica,
                "nivel": nivel,
            },
        )
        if created:
            self.stats["carrera"]["creados"] += 1
            return carrera

        updated_fields = []
        if carrera.nombre != nombre:
            carrera.nombre = nombre
            updated_fields.append("nombre")
        if carrera.unidad_academica_id != unidad_academica.id:
            carrera.unidad_academica = unidad_academica
            updated_fields.append("unidad_academica")
        if carrera.nivel != nivel:
            carrera.nivel = nivel
            updated_fields.append("nivel")

        if updated_fields:
            carrera.save(update_fields=updated_fields)
            self.stats["carrera"]["actualizados"] += 1
        else:
            self.stats["carrera"]["sin_cambios"] += 1
        return carrera

    def _upsert_plan_estudio(
        self,
        *,
        carrera: Carrera,
        nombre: str,
        descripcion: str,
        ordenanza: str,
        creditos: int,
        vigente_desde,
        vigente_hasta,
    ) -> PlanEstudio:
        plan, created = PlanEstudio.objects.get_or_create(
            ordenanza=ordenanza,
            defaults={
                "carrera": carrera,
                "nombre": nombre,
                "descripcion": descripcion,
                "creditos": creditos,
                "vigente_desde": vigente_desde,
                "vigente_hasta": vigente_hasta,
            },
        )
        if created:
            self.stats["plan_estudio"]["creados"] += 1
            return plan

        updated_fields = []
        field_values = {
            "carrera": carrera,
            "nombre": nombre,
            "descripcion": descripcion,
            "creditos": creditos,
            "vigente_desde": vigente_desde,
            "vigente_hasta": vigente_hasta,
        }
        for field_name, expected_value in field_values.items():
            current_value = getattr(plan, field_name)
            if current_value != expected_value:
                setattr(plan, field_name, expected_value)
                updated_fields.append(field_name)

        if updated_fields:
            plan.save(update_fields=updated_fields)
            self.stats["plan_estudio"]["actualizados"] += 1
        else:
            self.stats["plan_estudio"]["sin_cambios"] += 1
        return plan

    def _upsert_espacio_curricular(self, **data) -> EspacioCurricular:
        espacio, created = EspacioCurricular.objects.get_or_create(
            codigo=data["codigo"],
            defaults=data,
        )
        if created:
            self.stats["espacio_curricular"]["creados"] += 1
            return espacio

        updated_fields = []
        for field_name, expected_value in data.items():
            if field_name == "codigo":
                continue
            if getattr(espacio, field_name) != expected_value:
                setattr(espacio, field_name, expected_value)
                updated_fields.append(field_name)

        if updated_fields:
            espacio.save(update_fields=updated_fields)
            self.stats["espacio_curricular"]["actualizados"] += 1
        else:
            self.stats["espacio_curricular"]["sin_cambios"] += 1
        return espacio

    def _upsert_plan_estudio_ec(self, *, plan: PlanEstudio, espacio: EspacioCurricular) -> None:
        _, created = PlanEstudioEC.objects.get_or_create(plan_estudio=plan, espacio_curricular=espacio)
        if created:
            self.stats["plan_estudio_ec"]["creados"] += 1
        else:
            self.stats["plan_estudio_ec"]["sin_cambios"] += 1

    def _upsert_competencia(self, *, plan: PlanEstudio, codigo: str, nombre: str, descripcion: str) -> Competencia:
        competencia, created = Competencia.objects.get_or_create(
            plan_estudio=plan,
            codigo=codigo,
            defaults={
                "nombre": nombre,
                "descripcion": descripcion,
                "activo": True,
            },
        )
        if created:
            self.stats["competencia"]["creados"] += 1
            return competencia

        updated_fields = []
        if competencia.nombre != nombre:
            competencia.nombre = nombre
            updated_fields.append("nombre")
        if competencia.descripcion != descripcion:
            competencia.descripcion = descripcion
            updated_fields.append("descripcion")
        if competencia.activo is not True:
            competencia.activo = True
            updated_fields.append("activo")

        if updated_fields:
            competencia.save(update_fields=updated_fields)
            self.stats["competencia"]["actualizados"] += 1
        else:
            self.stats["competencia"]["sin_cambios"] += 1
        return competencia

    def _upsert_docente_user(self, row, row_number: int):
        email = self._clean_string(row.get("mail_en_siu")).lower()
        last_name = self._require_value(row, "apellido", row_number, "Docentes")
        first_name = self._require_value(row, "nombres", row_number, "Docentes")
        username_base = self._build_docente_username(row, email)

        user = None
        if email:
            user = User.objects.filter(email__iexact=email).first()
        if user is None:
            user = User.objects.filter(username=username_base).first()

        if user is None:
            username = self._ensure_unique_username(username_base)
            user = User(
                username=username,
                first_name=first_name,
                last_name=last_name,
                email=email,
                is_active=True,
            )
            user.set_unusable_password()
            user.save()
            self.stats["docente_user"]["creados"] += 1
            return user

        updated_fields = []
        if user.first_name != first_name:
            user.first_name = first_name
            updated_fields.append("first_name")
        if user.last_name != last_name:
            user.last_name = last_name
            updated_fields.append("last_name")
        if email and user.email.lower() != email:
            user.email = email
            updated_fields.append("email")
        if user.is_active is not True:
            user.is_active = True
            updated_fields.append("is_active")

        if updated_fields:
            user.save(update_fields=updated_fields)
            self.stats["docente_user"]["actualizados"] += 1
        else:
            self.stats["docente_user"]["sin_cambios"] += 1
        return user

    def _upsert_asignacion_docente(self, *, docente, espacio: EspacioCurricular, categoria: str, vigente_desde):
        asignacion = AsignacionDocente.objects.filter(
            docente=docente,
            espacio_curricular=espacio,
            vigente_desde=vigente_desde,
            vigente_hasta__isnull=True,
        ).first()
        if asignacion is None:
            asignacion = AsignacionDocente.objects.create(
                docente=docente,
                espacio_curricular=espacio,
                categoria=categoria,
                vigente_desde=vigente_desde,
                vigente_hasta=None,
            )
            self.stats["asignacion_docente"]["creados"] += 1
            return asignacion

        if asignacion.categoria != categoria:
            asignacion.categoria = categoria
            asignacion.save(update_fields=["categoria"])
            self.stats["asignacion_docente"]["actualizados"] += 1
        else:
            self.stats["asignacion_docente"]["sin_cambios"] += 1
        return asignacion

    def _print_summary(self):
        self.stdout.write("Resumen IMPORT_XLSX:")
        for label in [
            "unidad_academica",
            "carrera",
            "plan_estudio",
            "espacio_curricular",
            "plan_estudio_ec",
            "competencia",
            "docente_user",
            "asignacion_docente",
        ]:
            counters = self.stats[label]
            self.stdout.write(
                f"- {label}: creados={counters['creados']}, actualizados={counters['actualizados']}, sin_cambios={counters['sin_cambios']}"
            )

        for warning in self.warnings:
            self.stdout.write(self.style.WARNING(f"Aviso: {warning}"))

    def _clean_string(self, value) -> str:
        if value is None:
            return ""
        return str(value).strip()

    def _require_value(self, row, key: str, row_number: int, sheet_name: str) -> str:
        value = self._clean_string(row.get(key))
        if not value:
            raise CommandError(f"Hoja '{sheet_name}', fila {row_number}: la columna '{key}' está vacía.")
        return value

    def _normalize_code(self, value) -> str:
        if value is None:
            return ""
        if isinstance(value, int):
            return str(value)
        if isinstance(value, float):
            if value.is_integer():
                return str(int(value))
            return str(value).strip()

        text = str(value).strip()
        try:
            decimal = Decimal(text)
        except InvalidOperation:
            return text

        if decimal == decimal.to_integral_value():
            return str(int(decimal))
        return text

    def _extract_plan_aliases(self, raw_value: str) -> list[str]:
        aliases = []
        for chunk in raw_value.split(","):
            alias = self._normalize_code(chunk)
            if alias:
                aliases.append(alias)
        return aliases

    def _build_docente_username(self, row, email: str) -> str:
        for key in ["legajo", "id_docente"]:
            candidate = self._normalize_code(row.get(key))
            if candidate:
                return f"docente-{candidate}".lower()

        if email:
            return email.split("@", 1)[0].lower()

        full_name = f"{self._clean_string(row.get('apellido'))}-{self._clean_string(row.get('nombres'))}"
        slug = "".join(character.lower() if character.isalnum() else "-" for character in full_name)
        slug = "-".join(filter(None, slug.split("-")))
        if slug:
            return f"docente-{slug}"

        raise CommandError("No se pudo construir un username para un docente sin legajo, id_docente ni email.")

    def _ensure_unique_username(self, username_base: str) -> str:
        username = username_base[:150]
        suffix = 1
        while User.objects.filter(username=username).exists():
            tail = f"-{suffix}"
            username = f"{username_base[:150 - len(tail)]}{tail}"
            suffix += 1
        return username

    def _map_docente_categoria(self, raw_value, row_number: int) -> str:
        categoria = self._clean_string(raw_value).lower()
        if categoria == "titular":
            return AsignacionDocente.Categoria.TITULAR
        if categoria == "adjunto":
            return AsignacionDocente.Categoria.ADJUNTO
        if categoria == "asociado":
            return AsignacionDocente.Categoria.ASOCIADO
        if categoria == "jtp":
            return AsignacionDocente.Categoria.JTP
        if "ayudante" in categoria and "1" in categoria:
            return AsignacionDocente.Categoria.AYUDANTE_1
        if "ayudante" in categoria and "2" in categoria:
            return AsignacionDocente.Categoria.AYUDANTE_2
        raise CommandError(
            f"Hoja 'Docentes', fila {row_number}: categoría docente '{self._clean_string(raw_value)}' no mapeada."
        )

    def _to_int(self, value) -> int:
        if value in (None, ""):
            return 0
        if isinstance(value, int):
            return value
        if isinstance(value, float):
            return int(round(value))
        text = str(value).strip().replace(",", ".")
        if not text:
            return 0
        return int(round(float(text)))

    def _extract_year(self, value, row_number: int) -> int:
        text = self._clean_string(value)
        digits = "".join(character for character in text if character.isdigit())
        if not digits:
            raise CommandError(f"Hoja 'Espacios', fila {row_number}: no se pudo inferir el año de cursada desde '{text}'.")
        return int(digits)

    def _infer_periodo(self, cuatrimestre, bimestre) -> str:
        semester_text = self._clean_string(cuatrimestre).lower()
        bimester_text = self._clean_string(bimestre).lower()
        combined = f"{semester_text} {bimester_text}"
        if "anual" in combined:
            return EspacioCurricular.Periodo.ANUAL
        if "1" in semester_text or "1er" in combined:
            return EspacioCurricular.Periodo.PRIMER_SEMESTRE
        if "2" in semester_text or "2do" in combined:
            return EspacioCurricular.Periodo.SEGUNDO_SEMESTRE
        return EspacioCurricular.Periodo.ANUAL

    def _infer_tipo_espacio(self, horas_ip: int, horas_ta: int) -> str:
        total = horas_ip + horas_ta
        if total <= 0:
            return EspacioCurricular.TipoEspacio.T1_ASIGNATURAS

        ratio = horas_ip / total
        expected_ratios = {
            EspacioCurricular.TipoEspacio.T1_ASIGNATURAS: 0.30,
            EspacioCurricular.TipoEspacio.T2_SEMINARIOS: 0.50,
            EspacioCurricular.TipoEspacio.T3_TALLERES_LAB: 0.70,
            EspacioCurricular.TipoEspacio.T4_ACTIV_PROF_ESPEC: 0.80,
        }
        return min(expected_ratios, key=lambda choice: abs(expected_ratios[choice] - ratio))

    def _infer_end_date(self, estado_value):
        status = self._clean_string(estado_value).lower()
        if not status:
            return None
        if any(token in status for token in ["inactivo", "no vigente", "baja", "cerrado"]):
            return timezone.localdate()
        return None
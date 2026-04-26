from io import StringIO
from pathlib import Path
from tempfile import TemporaryDirectory

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.test import TestCase
from openpyxl import Workbook

from academics.models import Carrera, Competencia, EspacioCurricular, PlanEstudio, PlanEstudioEC, UnidadAcademica
from planning.models import AsignacionDocente


User = get_user_model()


class ImportAcademicXlsxCommandTests(TestCase):
	def _build_workbook(self, file_path: Path) -> None:
		workbook = Workbook()

		propuestas = workbook.active
		propuestas.title = "Propuestas"
		propuestas.append(["id_propuesta", "estado", "nombre", "nombre_plan", "titulo", "normativa"])
		propuestas.append([1, "Activo vigente", "CONTADOR PUBLICO", "CP26", "CP26, AUC26", "Ord. 110/2025- CS"])

		espacios = workbook.create_sheet("Espacios")
		espacios.append([
			"id_propuesta",
			"codigo_materia",
			"nombre_materia",
			"año_materia",
			"cuatrimestre_materia",
			"bimestre_materia",
			"correlativa",
			"formato",
			"caracter",
			"regimen",
			"area",
			"horas_i_p",
			"horas_f_p",
			"horas_t_a_e",
			"horas_semanales",
			"horas_total_estudiante",
			"creditos",
			"competencias_generales",
			"competencias_especificas",
			"resultado_aprendizaje",
			"contenido_minimo",
		])
		espacios.append([
			1,
			510201,
			"Derecho Público",
			"1°",
			"1er C",
			"1er B",
			"-",
			"Teórico",
			"Obligatoria",
			"Presencial",
			"J",
			70,
			20,
			80,
			20,
			150,
			6,
			"CG1",
			"CE1",
			"RA1",
			"Contenido",
		])

		competencias = workbook.create_sheet("Competencias")
		competencias.append(["titulo", "codigo", "descripcion", "tipo"])
		competencias.append([
			"CP26",
			"CE-G01",
			"Identificar y comprender situaciones problemáticas en contextos profesionales.",
			"Intelectuales",
		])

		docentes = workbook.create_sheet("Docentes")
		docentes.append([
			"nombre",
			"anio_de_cursada",
			"cuatrimestre",
			"codigo",
			"nombre-2",
			"id_docente",
			"legajo",
			"apellido",
			"nombres",
			"nro_documento",
			"nombre-3",
			"mail_en_siu",
		])
		docentes.append([
			"CONTADOR PUBLICO",
			1,
			"1er Cuatrimestre",
			510201,
			"Derecho Público",
			450,
			33591,
			"CASSAB",
			"Julieta Laura",
			32316496,
			"Titular",
			"docente@example.com",
		])

		workbook.save(file_path)

	def test_import_academic_xlsx_creates_core_entities(self):
		with TemporaryDirectory() as temp_dir:
			workbook_path = Path(temp_dir) / "import.xlsx"
			self._build_workbook(workbook_path)

			output = StringIO()
			call_command(
				"import_academic_xlsx",
				str(workbook_path),
				"--unidad-sigla=FCE",
				"--unidad-nombre=Facultad de Ciencias Económicas",
				"--career-level=G",
				"--plan-credits=180",
				"--vigente-desde=2025-01-01",
				stdout=output,
			)

		self.assertEqual(UnidadAcademica.objects.count(), 1)
		self.assertEqual(Carrera.objects.count(), 1)
		self.assertEqual(PlanEstudio.objects.count(), 1)
		self.assertEqual(EspacioCurricular.objects.count(), 1)
		self.assertEqual(PlanEstudioEC.objects.count(), 1)
		self.assertEqual(Competencia.objects.count(), 1)
		self.assertEqual(User.objects.count(), 1)
		self.assertEqual(AsignacionDocente.objects.count(), 1)

		carrera = Carrera.objects.get()
		self.assertEqual(carrera.codigo, "CP26")
		self.assertEqual(carrera.nombre, "CONTADOR PUBLICO")

		plan = PlanEstudio.objects.get()
		self.assertEqual(plan.nombre, "CP26")
		self.assertEqual(plan.descripcion, "CP26, AUC26")
		self.assertEqual(plan.ordenanza, "Ord. 110/2025- CS")
		self.assertEqual(plan.creditos, 180)

		espacio = EspacioCurricular.objects.get()
		self.assertEqual(espacio.codigo, "510201")
		self.assertEqual(espacio.horas_ip, 90)
		self.assertEqual(espacio.horas_ta, 80)
		self.assertEqual(espacio.creditos, 6)

		competencia = Competencia.objects.get()
		self.assertEqual(competencia.plan_estudio, plan)
		self.assertEqual(competencia.codigo, "CE-G01")
		self.assertIn("Intelectuales", competencia.nombre)

		docente = User.objects.get()
		self.assertEqual(docente.username, "docente-33591")
		self.assertEqual(docente.first_name, "Julieta Laura")
		self.assertEqual(docente.last_name, "CASSAB")
		self.assertEqual(docente.email, "docente@example.com")
		self.assertFalse(docente.has_usable_password())

		asignacion = AsignacionDocente.objects.get()
		self.assertEqual(asignacion.docente, docente)
		self.assertEqual(asignacion.espacio_curricular, espacio)
		self.assertEqual(asignacion.categoria, AsignacionDocente.Categoria.TITULAR)
		self.assertIn("docente_user", output.getvalue())

	def test_import_academic_xlsx_dry_run_rolls_back_changes(self):
		with TemporaryDirectory() as temp_dir:
			workbook_path = Path(temp_dir) / "import.xlsx"
			self._build_workbook(workbook_path)

			call_command(
				"import_academic_xlsx",
				str(workbook_path),
				"--unidad-sigla=FCE",
				"--unidad-nombre=Facultad de Ciencias Económicas",
				"--career-level=G",
				"--plan-credits=180",
				"--vigente-desde=2025-01-01",
				"--dry-run",
			)

		self.assertEqual(UnidadAcademica.objects.count(), 0)
		self.assertEqual(Carrera.objects.count(), 0)
		self.assertEqual(PlanEstudio.objects.count(), 0)
		self.assertEqual(EspacioCurricular.objects.count(), 0)
		self.assertEqual(PlanEstudioEC.objects.count(), 0)
		self.assertEqual(Competencia.objects.count(), 0)
		self.assertEqual(User.objects.count(), 0)
		self.assertEqual(AsignacionDocente.objects.count(), 0)

	def test_import_academic_xlsx_accepts_plan_aliases_in_competencias(self):
		with TemporaryDirectory() as temp_dir:
			workbook_path = Path(temp_dir) / "import_alias.xlsx"
			self._build_workbook(workbook_path)

			from openpyxl import load_workbook

			workbook = load_workbook(workbook_path)
			competencias = workbook["Competencias"]
			competencias.append([
				"AUC26",
				"CE-G02",
				"Aplicar criterios de análisis en contextos contables complejos.",
				"Intelectuales",
			])
			workbook.save(workbook_path)

			call_command(
				"import_academic_xlsx",
				str(workbook_path),
				"--unidad-sigla=FCE",
				"--unidad-nombre=Facultad de Ciencias Económicas",
				"--career-level=G",
				"--plan-credits=180",
				"--vigente-desde=2025-01-01",
			)

		self.assertEqual(Competencia.objects.count(), 2)
		plan = PlanEstudio.objects.get()
		alias_competencia = Competencia.objects.get(codigo="CE-G02")
		self.assertEqual(alias_competencia.plan_estudio, plan)

	def test_import_academic_xlsx_is_idempotent_for_docentes(self):
		with TemporaryDirectory() as temp_dir:
			workbook_path = Path(temp_dir) / "import_idempotent.xlsx"
			self._build_workbook(workbook_path)

			for _ in range(2):
				call_command(
					"import_academic_xlsx",
					str(workbook_path),
					"--unidad-sigla=FCE",
					"--unidad-nombre=Facultad de Ciencias Económicas",
					"--career-level=G",
					"--plan-credits=180",
					"--vigente-desde=2025-01-01",
				)

		self.assertEqual(User.objects.count(), 1)
		self.assertEqual(AsignacionDocente.objects.count(), 1)

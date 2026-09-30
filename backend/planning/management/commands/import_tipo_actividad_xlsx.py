from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from planning.models import TipoActividad


class DryRunRollback(Exception):
    """Internal exception to rollback dry-run imports."""


class Command(BaseCommand):
    help = "Importa tipificaciones de actividades pedagogicas desde un XLSX."

    def add_arguments(self, parser):
        parser.add_argument("xlsx_path", help="Ruta al archivo XLSX de tipificaciones")
        parser.add_argument(
            "--tipo-dedicacion",
            default=TipoActividad.TipoDedicacion.TRABAJO_AUTONOMO,
            choices=[choice for choice, _ in TipoActividad.TipoDedicacion.choices],
            help="Tipo de dedicacion a asignar a las tipificaciones importadas",
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Procesa el archivo sin persistir cambios",
        )

    def handle(self, *args, **options):
        workbook_path = Path(options["xlsx_path"]).expanduser().resolve()
        if not workbook_path.exists():
            raise CommandError(f"No existe el archivo: {workbook_path}")

        workbook = load_workbook(filename=workbook_path, data_only=True)
        worksheet = workbook[workbook.sheetnames[0]]
        self.stats = defaultdict(lambda: {"creados": 0, "actualizados": 0, "sin_cambios": 0})

        try:
            with transaction.atomic():
                self._import_tipificaciones(
                    worksheet,
                    tipo_dedicacion=options["tipo_dedicacion"],
                )
                if options["dry_run"]:
                    raise DryRunRollback()
        except DryRunRollback:
            self.stdout.write(self.style.WARNING("Dry-run ejecutado: los cambios fueron revertidos."))

        self._print_summary()

    def _import_tipificaciones(self, worksheet, *, tipo_dedicacion: str) -> None:
        if worksheet.max_row < 3:
            raise CommandError("La planilla no contiene filas de tipificaciones para importar.")

        for row_number in range(3, worksheet.max_row + 1):
            nombre = self._clean_string(worksheet[f"A{row_number}"].value)
            if not nombre:
                continue

            descripcion_base = self._clean_string(worksheet[f"B{row_number}"].value)
            ejemplos = self._clean_string(worksheet[f"C{row_number}"].value)
            notas = self._clean_string(worksheet[f"F{row_number}"].value)

            descripcion = self._compose_descripcion(
                descripcion=descripcion_base,
                ejemplos=ejemplos,
                notas=notas,
            )
            self._upsert_tipo_actividad(
                nombre=nombre,
                descripcion=descripcion,
                tipo_dedicacion=tipo_dedicacion,
            )

    def _upsert_tipo_actividad(self, *, nombre: str, descripcion: str, tipo_dedicacion: str):
        tipo_actividad, created = TipoActividad.objects.get_or_create(
            nombre=nombre,
            defaults={
                "descripcion": descripcion,
                "tipo_dedicacion": tipo_dedicacion,
            },
        )
        if created:
            self.stats["tipo_actividad"]["creados"] += 1
            return tipo_actividad

        if tipo_actividad.tipo_dedicacion != tipo_dedicacion:
            raise CommandError(
                f'La tipificación "{nombre}" ya existe con dedicación '
                f'"{tipo_actividad.tipo_dedicacion}"; no se cambiará a "{tipo_dedicacion}".'
            )

        updated_fields = []
        if tipo_actividad.descripcion != descripcion:
            tipo_actividad.descripcion = descripcion
            updated_fields.append("descripcion")

        if updated_fields:
            tipo_actividad.save(update_fields=updated_fields)
            self.stats["tipo_actividad"]["actualizados"] += 1
        else:
            self.stats["tipo_actividad"]["sin_cambios"] += 1
        return tipo_actividad

    def _compose_descripcion(self, *, descripcion: str, ejemplos: str, notas: str) -> str:
        chunks = []
        if descripcion:
            chunks.append(descripcion)
        if ejemplos:
            chunks.append(f"Ejemplos: {ejemplos}")
        if notas:
            chunks.append(f"Consideraciones: {notas}")
        return "\n\n".join(chunks)

    def _print_summary(self) -> None:
        counters = self.stats["tipo_actividad"]
        self.stdout.write("Resumen IMPORT_TIPO_ACTIVIDAD_XLSX:")
        self.stdout.write(
            f"- tipo_actividad: creados={counters['creados']}, actualizados={counters['actualizados']}, sin_cambios={counters['sin_cambios']}"
        )

    def _clean_string(self, value) -> str:
        if value is None:
            return ""
        return str(value).strip()
